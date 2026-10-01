import re, json, io, urllib.request
H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
     "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8", "Accept-Language": "en-US,en;q=0.9",
     "Sec-Fetch-Dest": "document", "Sec-Fetch-Mode": "navigate", "Sec-Fetch-Site": "none", "Upgrade-Insecure-Requests": "1"}
def get(url):
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers=H), timeout=40) as r:
            b = r.read(); print("STATUS", r.status, r.headers.get("content-type"), len(b)); return b
    except Exception as e:
        print("ERR", type(e).__name__, str(e)[:150]); return None
def txt(b):
    t = b.decode("utf-8", "replace"); t = re.sub(r"<script.*?</script>|<style.*?</style>", " ", t, flags=re.S)
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", t))
for name, url in [("ici", "https://www.ici.org/research/stats/flows"), ("ici_combined", "https://www.ici.org/research/stats/combined_flows")]:
    print("\n####", name); b = get(url)
    if b:
        s = b.decode("utf-8","replace"); t = txt(b)
        print("LINKS", sorted(set(re.findall(r'href="([^"]+\.(?:xlsx?|csv|pdf))"', s)))[:15])
        i = t.find("Estimated"); print("CTX", t[i:i+1500])
print("\n#### stax links")
b = get("https://www.schwab.com/investment-research/stax")
if b:
    s = b.decode("utf-8","replace")
    for m in re.finditer(r'<a[^>]+href="([^"]+)"[^>]*>\s*(?:<[^>]+>\s*)*Explore the STAX', s): print("EXPLORE", m.group(1))
    print("ALL", [u for u in sorted(set(re.findall(r'href="([^"]+)"', s))) if "stax" in u.lower()][:20])
print("\n#### jpx weekly xlsx")
b = get("https://www.jpx.co.jp/english/markets/statistics-equities/margin/dreu2500000073tg-att/20260925_mtcurrent.xlsx")
if b:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
    for ws in wb.worksheets[:2]:
        print("SHEET", ws.title, ws.max_row, ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=45, values_only=True):
            if any(c is not None for c in row): print(" ", [c for c in row if c is not None][:14])
